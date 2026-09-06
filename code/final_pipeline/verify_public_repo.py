# -*- coding: utf-8 -*-
"""Check that the public repository contains the final release materials."""

from pathlib import Path
import hashlib
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]

REQUIRED_SCRIPTS = {
    "build_final_supplementary_tables.py",
    "make_prisma_integrated_figure.py",
    "p01_full_figures.py",
    "p01_full_rerun.py",
    "p13_classification_sensitivity.py",
    "p15_mixed_covariate_diagnostics.py",
    "p16_ocr_sensitivity.py",
    "p25_gained_reporting_table.py",
    "prepare_revised_figures.py",
    "recompute_final_classification_stability.py",
    "recompute_formal_deconv_final.py",
    "run_consensus_mibc.R",
    "run_p01_regmedint_survival.R",
}

REQUIRED_FIGURES = [
    "Figure1",
    "Figure2",
    "FigureS1",
    "FigureS2",
    "FigureS3",
    "FigureS4",
    "FigureS5",
    "FigureS6",
]

REQUIRED_DATA = [
    "data/processed/p01_full/final_classification_bootstrap_stability.csv",
    "data/processed/p01_full/final_signature_sensitivity_formal_deconvolution.csv",
    "data/processed/p01_full/p25_gained_reporting_table.csv",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    checks = []
    for script in REQUIRED_SCRIPTS:
        path = ROOT / "code" / "final_pipeline" / script
        checks.append((f"script {script}", path.exists()))

    for name in REQUIRED_FIGURES:
        checks.append((f"figure {name}.png", (ROOT / "figures" / f"{name}.png").exists()))
        checks.append((f"figure {name}.pdf", (ROOT / "figures" / f"{name}.pdf").exists()))

    for relative in REQUIRED_DATA:
        checks.append((f"data {relative}", (ROOT / relative).exists()))

    workbook = ROOT / "supplementary" / "Supplementary_Tables_S1_S72_UroOnc_revised.xlsx"
    if workbook.exists():
        sheets = load_workbook(workbook, read_only=True).sheetnames
        numbers = {
            int(re.match(r"S(\d+)", sheet).group(1))
            for sheet in sheets
            if re.match(r"S\d+", sheet)
        }
        checks.append(("S1-S72 workbook complete", numbers == set(range(1, 73))))
    else:
        checks.append(("S1-S72 workbook exists", False))

    checks.append(("MIT license", (ROOT / "LICENSE-CODE").exists()))
    checks.append(("CC BY 4.0 license", (ROOT / "LICENSE-DATA").exists()))
    checks.append(
        ("scikit-learn in requirements", "scikit-learn" in (ROOT / "requirements.txt").read_text())
    )

    code_files = list((ROOT / "code" / "final_pipeline").rglob("*"))
    absolute_paths = []
    for path in code_files:
        if path.name == "verify_public_repo.py":
            continue
        if path.suffix.lower() not in {".py", ".r"}:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        if re.search(r"C:/Users|C:\\Users|setwd", text):
            absolute_paths.append(path.name)
    checks.append(("no absolute paths in final_pipeline", not absolute_paths))

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    checks.append(
        ("README references final workbook",
         "Supplementary_Tables_S1_S72_UroOnc_revised.xlsx" in readme)
    )
    checks.append(("README references licenses", "LICENSE-CODE" in readme and "LICENSE-DATA" in readme))
    checks.append(
        ("stale helper removed", not (ROOT / "sync_final_release_files.py").exists())
    )

    failed = [name for name, passed in checks if not passed]
    for name, passed in checks:
        print(("PASS" if passed else "FAIL"), name)
    print("FAILED", failed)
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
